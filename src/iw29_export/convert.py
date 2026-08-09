"""Turn whatever SAP wrote into a clean, typed .xlsx and tidy in-memory rows.

Letting SAP produce text and building the workbook here avoids the two classic
problems with SAP's own Excel export: it needs Excel installed and visible, and
it hands back strings where you wanted dates and numbers.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, List, Optional, Sequence, Tuple

from . import xlsx
from .errors import ExportError
from .logging_setup import get_logger

log = get_logger("convert")

_ENCODING_CANDIDATES = ("utf-8-sig", "utf-16", "cp1252", "latin-1")
_DATE_FORMATS = ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y%m%d")
# YYYYMMDD must stay inside a real calendar window. Without this, SAP order
# numbers such as 73231022 become date(7323, 10, 22) and break XLOOKUP keys.
_YYYYMMDD_YEAR_MIN = 1990
_YYYYMMDD_YEAR_MAX = 2100
_EXCEL_EPOCH = date(1899, 12, 30)
_SEPARATOR_LINE = re.compile(r"^[\s|+\-=_]*$")
_NUMBER_LIKE = re.compile(r"^-?[\d.,\s]+-?$")
_SEPARATORS = ("\t", "|", ";")
_PREAMBLE_LIMIT = 30

# Identifier columns must remain text so Excel never re-interprets them as dates.
_TEXT_ID_HEADERS = frozenset(
    {
        "order",
        "notification",
        "message",
        "functional location",
        "sort field",
        "revision",
        "mn.wk.ctr",
    }
)
_DATE_HEADERS = frozenset(
    {
        "basic fin.",
        "bsc start",
        "actual end",
        "created on",
        "schedstart",
        "sched start",
        "schedfinish",
        "sched finish",
    }
)


@dataclass
class Table:
    headers: List[str]
    rows: List[List[Any]]

    @property
    def row_count(self) -> int:
        return len(self.rows)


def read_sap_text(path: Path) -> Table:
    """Read a SAP list export, whether it is tab delimited or pipe framed."""
    text, encoding = _read_text(path)
    log.info("Read %s as %s (%d characters)", path.name, encoding, len(text))

    lines = [line for line in text.splitlines() if not _SEPARATOR_LINE.match(line)]
    if not lines:
        raise ExportError(f"{path.name} contains no usable rows.")

    separator, header_index = _find_header(lines, path.name)
    if header_index:
        log.info(
            "Skipped %d preamble line(s) before the header, e.g. %r",
            header_index,
            lines[0][:60].strip(),
        )

    raw_rows = [_split(line, separator) for line in lines[header_index:]]
    header_cells = raw_rows[0]
    drop_first = _has_empty_lead_column(raw_rows)
    if drop_first:
        raw_rows = [row[1:] for row in raw_rows]
        header_cells = raw_rows[0]

    headers = _clean_headers(header_cells)
    width = len(headers)
    header_signature = [cell.strip() for cell in header_cells]

    rows: List[List[Any]] = []
    for raw in raw_rows[1:]:
        cells = [str(cell).strip() for cell in raw[:width]]
        cells += [""] * (width - len(cells))
        if not any(cells):
            continue
        # Long lists repeat the page title and header at each page break.
        if cells == header_signature:
            continue
        rows.append(
            [
                coerce(cell, header=headers[index] if index < len(headers) else "")
                for index, cell in enumerate(cells)
            ]
        )

    if not rows:
        raise ExportError(f"{path.name} has a header but no data rows.")
    log.info("Parsed %d data rows across %d columns", len(rows), width)
    return normalise_identifier_columns(Table(headers=headers, rows=rows))


def _find_header(lines: Sequence[str], name: str) -> Tuple[str, int]:
    """Locate the header row and its separator.

    SAP puts a page title and blank lines above the data, so the separator
    cannot be read off the first line. The header is the first line that is
    actually split into several columns.
    """
    window = lines[:_PREAMBLE_LIMIT]
    totals = {sep: sum(line.count(sep) for line in window) for sep in _SEPARATORS}
    separator = max(_SEPARATORS, key=lambda candidate: totals[candidate])
    if totals[separator]:
        for index, line in enumerate(window):
            if separator in line:
                return separator, index
    raise ExportError(
        f"Could not work out the column separator in {name}. Expected tab, pipe or "
        f"semicolon in the first {_PREAMBLE_LIMIT} lines. If SAP wrote a fixed-width "
        "list, re-run so the 'Text with Tabs' format is chosen."
    )


def _split(line: str, separator: str) -> List[str]:
    if separator == "|":
        return [cell.strip() for cell in line.strip().strip("|").split("|")]
    return next(csv.reader([line], delimiter=separator))


def _has_empty_lead_column(raw_rows: Sequence[Sequence[str]]) -> bool:
    """SAP's tab export starts every line with a tab, giving a blank first column."""
    sample = raw_rows[: min(len(raw_rows), 50)]
    return len(sample) > 1 and all(
        row and not str(row[0]).strip() for row in sample
    )


def read_xlsx(path: Path) -> Table:
    """Read back a workbook SAP (or this app) produced.

    Prefers openpyxl when available; otherwise uses a stdlib ZIP/XML reader so
    corporate machines without PyPI can still repair / re-enrich harvests.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        return read_xlsx_stdlib(path)

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if header_row is None:
            raise ExportError(f"{path.name} has no rows.")
        headers = _clean_headers([str(c) if c is not None else "" for c in header_row])
        width = len(headers)
        rows = []
        for raw in iterator:
            cells = list(raw[:width]) + [None] * max(0, width - len(raw))
            if all(cell in (None, "") for cell in cells):
                continue
            typed: List[Any] = []
            for index, cell in enumerate(cells):
                header = headers[index] if index < len(headers) else ""
                if _is_text_id_header(header):
                    typed.append(identifier_text(cell))
                elif isinstance(cell, str):
                    typed.append(coerce(cell, header=header))
                else:
                    typed.append(cell)
            rows.append(typed)
    finally:
        workbook.close()

    if not rows:
        raise ExportError(f"{path.name} has a header but no data rows.")
    return normalise_identifier_columns(Table(headers=headers, rows=rows))


def read_xlsx_stdlib(path: Path) -> Table:
    """Stdlib .xlsx reader (first sheet) used when openpyxl is unavailable."""
    import zipfile
    from xml.etree import ElementTree as ET

    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

    def col_index(cell_ref: str) -> int:
        match = re.match(r"([A-Z]+)", cell_ref or "")
        if not match:
            return 0
        number = 0
        for char in match.group(1):
            number = number * 26 + (ord(char) - 64)
        return number

    with zipfile.ZipFile(path) as archive:
        strings: List[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for si in root.findall("m:si", ns):
                texts = [
                    node.text or ""
                    for node in si.iter(
                        "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t"
                    )
                ]
                strings.append("".join(texts))
        sheet_name = next(
            (
                name
                for name in archive.namelist()
                if name.startswith("xl/worksheets/sheet")
            ),
            None,
        )
        if not sheet_name:
            raise ExportError(f"{path.name} has no worksheet.")
        sheet = ET.fromstring(archive.read(sheet_name))

        def cell_val(cell: Any) -> Any:
            kind = cell.attrib.get("t")
            node = cell.find("m:v", ns)
            inline = cell.find("m:is", ns)
            if kind == "inlineStr" and inline is not None:
                return "".join(
                    node.text or ""
                    for node in inline.iter(
                        "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t"
                    )
                )
            if node is None:
                return None
            if kind == "s":
                index = int(node.text)
                return strings[index] if 0 <= index < len(strings) else None
            text = node.text
            try:
                number = float(text)
            except (TypeError, ValueError):
                return text
            if number.is_integer():
                return int(number)
            return number

        parsed: List[Dict[int, Any]] = []
        max_col = 0
        for row in sheet.findall("m:sheetData/m:row", ns):
            values: Dict[int, Any] = {}
            for cell in row.findall("m:c", ns):
                index = col_index(cell.attrib.get("r", ""))
                if not index:
                    continue
                values[index] = cell_val(cell)
                max_col = max(max_col, index)
            if values:
                parsed.append(values)

    if not parsed:
        raise ExportError(f"{path.name} has no rows.")
    width = max_col
    raw_headers = [parsed[0].get(index) for index in range(1, width + 1)]
    headers = _clean_headers(
        [str(cell) if cell is not None else "" for cell in raw_headers]
    )
    rows: List[List[Any]] = []
    for values in parsed[1:]:
        cells = [values.get(index) for index in range(1, width + 1)]
        if all(cell in (None, "") for cell in cells):
            continue
        typed = []
        for index, cell in enumerate(cells):
            header = headers[index] if index < len(headers) else ""
            if _is_text_id_header(header):
                typed.append(identifier_text(cell))
            elif _is_date_header(header):
                typed.append(excel_serial_to_date(cell) if not isinstance(cell, str) else coerce(cell, header=header))
            elif isinstance(cell, str):
                typed.append(coerce(cell, header=header))
            else:
                typed.append(cell)
        rows.append(typed)
    if not rows:
        raise ExportError(f"{path.name} has a header but no data rows.")
    return normalise_identifier_columns(Table(headers=headers, rows=rows))


def write_xlsx(table: Table, destination: Path, sheet_name: str = "Sheet1") -> Path:
    """Write a styled single-sheet workbook using only the standard library."""
    xlsx.write(
        headers=table.headers,
        rows=table.rows,
        destination=destination,
        sheet_name=sheet_name,
        column_widths=_column_widths(table),
    )
    log.info(
        "Workbook written: %s (%d rows, %d columns)",
        destination,
        table.row_count,
        len(table.headers),
    )
    return destination


def coerce(value: str, header: str = "") -> Any:
    """Best-effort conversion of a SAP cell into a date, number or trimmed string."""
    text = (value or "").strip()
    if not text or text in {"-", "--"}:
        return None

    if _is_text_id_header(header):
        return identifier_text(text)

    for fmt in _DATE_FORMATS:
        try:
            parsed = datetime.strptime(text, fmt)
        except ValueError:
            continue
        if fmt == "%Y%m%d":
            if not (_YYYYMMDD_YEAR_MIN <= parsed.year <= _YYYYMMDD_YEAR_MAX):
                continue
        elif parsed.year <= 1900:
            continue
        return parsed.date()

    # Leading zeros are meaningful in SAP keys, so those stay text.
    if _NUMBER_LIKE.match(text) and not (len(text) > 1 and text.startswith("0")):
        number = _parse_number(text)
        if number is not None:
            return number
    return text


def identifier_text(value: Any) -> str:
    """Canonical text form for Order / Notification-style identifiers.

    Repairs the classic Excel damage where an 8-digit order such as ``73231022``
    was parsed as ``date(7323, 10, 22)`` and later shown as ``7323-10-22``.
    """
    if value is None:
        return ""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return f"{value.year:04d}{value.month:02d}{value.day:02d}"
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        try:
            numeric = float(value)
        except (TypeError, ValueError):
            return str(value).strip()
        if not numeric.is_integer():
            return str(value).strip()
        serial = int(numeric)
        # Excel date serial left after a mangled YYYYMMDD parse (e.g. 1981006).
        # Real WO numbers are ~8 digits (>= 1e7) and must stay untouched.
        if 300_000 <= serial <= 3_000_000:
            decoded = _EXCEL_EPOCH + timedelta(days=serial)
            if decoded.year > _YYYYMMDD_YEAR_MAX or decoded.year < _YYYYMMDD_YEAR_MIN:
                return f"{decoded.year:04d}{decoded.month:02d}{decoded.day:02d}"
        return str(serial)

    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    match = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})$", text)
    if match:
        year, month, day = (
            int(match.group(1)),
            int(match.group(2)),
            int(match.group(3)),
        )
        if year > _YYYYMMDD_YEAR_MAX or year < _YYYYMMDD_YEAR_MIN:
            return f"{year:04d}{month:02d}{day:02d}"
    if text.isdigit():
        return str(int(text))
    return text


def normalise_identifier_columns(table: Table) -> Table:
    """Force identifier columns to stable text values (Order, Message, …)."""
    indexes = [
        index
        for index, header in enumerate(table.headers)
        if _is_text_id_header(header)
    ]
    if not indexes:
        return table
    rows: List[List[Any]] = []
    for row in table.rows:
        values = list(row)
        while len(values) < len(table.headers):
            values.append(None)
        for index in indexes:
            values[index] = identifier_text(values[index])
        rows.append(values)
    return Table(headers=list(table.headers), rows=rows)


def excel_serial_to_date(value: Any) -> Any:
    """Convert an Excel day-serial to ``date`` when it looks like a real date."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool) or value is None or value == "":
        return value
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return value
    if not numeric.is_integer():
        return value
    serial = int(numeric)
    # Excel dates used in maintenance plans sit roughly in 1955–2119.
    if 20_000 <= serial <= 80_000:
        return _EXCEL_EPOCH + timedelta(days=serial)
    return value


def _is_text_id_header(header: str) -> bool:
    name = re.sub(r"\s+", " ", str(header or "").strip()).lower()
    name = re.sub(r"\s+\(\d+\)$", "", name)
    return name in _TEXT_ID_HEADERS


def _is_date_header(header: str) -> bool:
    name = re.sub(r"\s+", " ", str(header or "").strip()).lower()
    name = re.sub(r"\s+\(\d+\)$", "", name)
    return name in _DATE_HEADERS


def _parse_number(text: str) -> Optional[Any]:
    cleaned = text.replace("\u00a0", "").replace(" ", "")
    negative = cleaned.endswith("-")
    cleaned = cleaned.rstrip("-")
    if not cleaned or not any(ch.isdigit() for ch in cleaned):
        return None

    # SAP formats numbers per user settings: 1.234,56 or 1,234.56.
    if "," in cleaned and "." in cleaned:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        decimals = cleaned.rsplit(",", 1)[1]
        cleaned = cleaned.replace(",", "." if len(decimals) != 3 else "")
    elif cleaned.count(".") == 1 and len(cleaned.rsplit(".", 1)[1]) == 3:
        cleaned = cleaned.replace(".", "")

    try:
        number: Any = float(cleaned)
    except ValueError:
        return None
    if negative:
        number = -number
    if number.is_integer() and "." not in cleaned:
        return int(number)
    return number


def _read_text(path: Path) -> Tuple[str, str]:
    data = path.read_bytes()
    if not data:
        raise ExportError(f"{path} is empty.")
    for encoding in _ENCODING_CANDIDATES:
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ExportError(f"Could not decode {path} with any of {_ENCODING_CANDIDATES}.")


def _clean_headers(raw: Sequence[str]) -> List[str]:
    headers: List[str] = []
    seen: dict = {}
    for index, cell in enumerate(raw, start=1):
        name = re.sub(r"\s+", " ", str(cell).strip()) or f"Column {index}"
        count = seen.get(name.lower(), 0) + 1
        seen[name.lower()] = count
        headers.append(name if count == 1 else f"{name} ({count})")
    return headers


def _column_widths(table: Table, minimum: int = 10, maximum: int = 55) -> List[int]:
    widths = []
    for index, header in enumerate(table.headers):
        longest = len(str(header))
        for row in table.rows[:500]:
            cell = row[index] if index < len(row) else None
            longest = max(longest, len(_display_length(cell)))
        widths.append(max(minimum, min(maximum, longest + 2)))
    return widths


def _display_length(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return "0000-00-00"
    return str(value)
